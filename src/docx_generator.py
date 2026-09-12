"""Generate a professional Affidavit-in-Reply DOCX from mapped content."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn

from src.content_mapper import MappedAffidavitContent
from src.schemas import (
    AffidavitSection,
    AffidavitTemplateSchema,
    OptionalTemplateElement,
    default_bombay_hc_affidavit_in_reply_template,
)


def _set_cell_margins(cell, top=0, start=0, bottom=0, end=0):
    """Set cell margins in twips."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.find(qn("w:tcMar"))
    if tcMar is None:
        from lxml import etree
        tcMar = etree.SubElement(tcPr, qn("w:tcMar"))
    for attr, val in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        el = tcMar.find(qn(f"w:{attr}"))
        if el is None:
            from lxml import etree
            el = etree.SubElement(tcMar, qn(f"w:{attr}"))
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")


class AffidavitDocxGenerator:
    """Render a MappedAffidavitContent into a .docx file using a reusable template schema."""

    def generate(
        self,
        content: MappedAffidavitContent,
        output_path: str | Path,
        template: AffidavitTemplateSchema | None = None,
    ) -> Path:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        template = template or content.template or default_bombay_hc_affidavit_in_reply_template()
        doc = Document()

        # --- Page setup ---
        section = doc.sections[0]
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)

        # --- Default font ---
        style = doc.styles["Normal"]
        font = style.font
        font.name = "Times New Roman"
        font.size = Pt(12)
        font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.line_spacing = 1.15

        # Render sections dynamically based on template schema
        for sec_spec in sorted(template.sections, key=lambda s: s.order):
            sec = sec_spec.section
            if sec == AffidavitSection.FORUM_HEADING:
                self._render_heading_section(doc, content.forum_heading, sec, template)
            elif sec == AffidavitSection.JURISDICTION:
                self._render_heading_section(doc, content.jurisdiction, sec, template)
            elif sec == AffidavitSection.CASE_NUMBER:
                self._render_heading_section(doc, content.case_number_line, sec, template)
                doc.add_paragraph()  # spacer
            elif sec == AffidavitSection.CAUSE_TITLE:
                self._add_cause_title(doc, content)
            elif sec == AffidavitSection.AFFIDAVIT_TITLE:
                doc.add_paragraph()
                self._render_heading_section(doc, content.affidavit_title, sec, template)
                doc.add_paragraph()
            elif sec == AffidavitSection.DEPONENT_CLAUSE:
                p = doc.add_paragraph()
                p.alignment = (
                    WD_ALIGN_PARAGRAPH.JUSTIFY
                    if template.formatting.paragraph_text_justified
                    else WD_ALIGN_PARAGRAPH.LEFT
                )
                p.add_run(content.deponent_clause.text)
                doc.add_paragraph()
            elif sec == AffidavitSection.NUMBERED_PARAGRAPHS:
                self._render_numbered_paragraphs(doc, content, template)
                doc.add_paragraph()
            elif sec == AffidavitSection.PRAYER:
                self._render_prayer(doc, content, template)
                doc.add_paragraph()
                doc.add_paragraph()
            elif sec == AffidavitSection.JURAT:
                self._render_jurat(doc, content, template)
                doc.add_paragraph()
            elif sec == AffidavitSection.VERIFICATION:
                self._render_verification(doc, content, template)
                doc.add_paragraph()

        # Optional elements (e.g. advocate block)
        if (
            OptionalTemplateElement.ADVOCATE_BLOCK in template.optional_elements
            and content.advocate
        ):
            self._render_advocate_block(doc, content.advocate)

        doc.save(str(output_path))
        return output_path

    def _render_heading_section(
        self,
        doc: Document,
        text: str,
        sec: AffidavitSection,
        template: AffidavitTemplateSchema,
    ) -> None:
        p = doc.add_paragraph()
        bold_caps_centred = sec in template.formatting.bold_caps_centred_sections
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold_caps_centred else WD_ALIGN_PARAGRAPH.LEFT
        run_text = text.upper() if bold_caps_centred else text
        run = p.add_run(run_text)
        run.bold = bold_caps_centred
        run.font.size = Pt(12)

    def _render_numbered_paragraphs(
        self,
        doc: Document,
        content: MappedAffidavitContent,
        template: AffidavitTemplateSchema,
    ) -> None:
        rule = template.paragraph_numbering
        suffix = rule.suffix if rule.suffix else "."
        for para in content.body_paragraphs:
            p = doc.add_paragraph()
            p.alignment = (
                WD_ALIGN_PARAGRAPH.JUSTIFY
                if template.formatting.paragraph_text_justified
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            run_num = p.add_run(f"{para.number}{suffix} ")
            run_num.bold = rule.number_bold
            p.add_run(para.text)
            p.paragraph_format.space_after = Pt(6)

    def _render_prayer(
        self,
        doc: Document,
        content: MappedAffidavitContent,
        template: AffidavitTemplateSchema,
    ) -> None:
        rule = template.prayer_lettering
        heading_text = rule.heading or content.prayer.heading
        self._add_heading_line(doc, heading_text)

        p = doc.add_paragraph()
        p.alignment = (
            WD_ALIGN_PARAGRAPH.JUSTIFY
            if template.formatting.paragraph_text_justified
            else WD_ALIGN_PARAGRAPH.LEFT
        )
        p.add_run(content.prayer.introductory_phrase)
        p.paragraph_format.space_after = Pt(6)

        suffix = rule.letter_suffix if rule.letter_suffix else ")"
        for clause in content.prayer.clauses:
            p = doc.add_paragraph()
            p.alignment = (
                WD_ALIGN_PARAGRAPH.JUSTIFY
                if template.formatting.paragraph_text_justified
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            p.paragraph_format.left_indent = Inches(0.5)
            prefix = f"({clause.letter}) " if suffix == ")" else f"{clause.letter}{suffix} "
            label_run = p.add_run(prefix)
            label_run.bold = rule.letters_bold
            p.add_run(clause.text)
            p.paragraph_format.space_after = Pt(4)

    def _render_jurat(
        self,
        doc: Document,
        content: MappedAffidavitContent,
        template: AffidavitTemplateSchema,
    ) -> None:
        if template.formatting.deponent_caps_right_aligned:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run("DEPONENT")
            run.bold = True
            doc.add_paragraph()

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.add_run(content.jurat.text)
        doc.add_paragraph()

        if template.formatting.before_me_left_aligned:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run("Before Me")
            run.bold = True

    def _render_verification(
        self,
        doc: Document,
        content: MappedAffidavitContent,
        template: AffidavitTemplateSchema,
    ) -> None:
        self._add_heading_line(doc, "VERIFICATION")
        doc.add_paragraph()
        p = doc.add_paragraph()
        p.alignment = (
            WD_ALIGN_PARAGRAPH.JUSTIFY
            if template.formatting.paragraph_text_justified
            else WD_ALIGN_PARAGRAPH.LEFT
        )
        p.add_run(content.verification.text)

    def _render_advocate_block(self, doc: Document, advocate) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.add_run(advocate.firm).bold = True
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.add_run(f"Advocates for {advocate.acting_for}")

    def _add_heading_line(self, doc: Document, text: str) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(12)

    def _add_cause_title(self, doc: Document, content: MappedAffidavitContent) -> None:
        ct = content.cause_title

        # Petitioner line
        table = doc.add_table(rows=1, cols=2)
        table.autofit = True
        table.columns[0].width = Inches(4.5)
        table.columns[1].width = Inches(1.5)

        row = table.rows[0]
        cell_name = row.cells[0]
        cell_tag = row.cells[1]

        cell_name.text = ct.petitioner.name
        cell_tag.text = ct.petitioner.status_tag

        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
        cell_tag.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

        # Remove table borders
        self._remove_table_borders(table)

        # VERSUS
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("VERSUS")
        run.bold = True
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)

        # Respondents
        for respondent in ct.respondents:
            table = doc.add_table(rows=1, cols=2)
            table.autofit = True
            table.columns[0].width = Inches(4.5)
            table.columns[1].width = Inches(1.5)

            row = table.rows[0]
            row.cells[0].text = respondent.name
            row.cells[1].text = respondent.status_tag

            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.space_before = Pt(0)
            row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

            self._remove_table_borders(table)

    def _remove_table_borders(self, table) -> None:
        """Remove all borders from a table for clean layout."""
        tbl = table._tbl
        tblPr = tbl.tblPr if tbl.tblPr is not None else tbl._add_tblPr()
        from lxml import etree
        borders = etree.SubElement(tblPr, qn("w:tblBorders"))
        for border_name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            border = etree.SubElement(borders, qn(f"w:{border_name}"))
            border.set(qn("w:val"), "none")
            border.set(qn("w:sz"), "0")
            border.set(qn("w:space"), "0")
            border.set(qn("w:color"), "auto")


def generate_affidavit_docx(
    content: MappedAffidavitContent,
    output_path: str | Path,
) -> Path:
    """Convenience wrapper around AffidavitDocxGenerator.generate."""
    return AffidavitDocxGenerator().generate(content, output_path)
